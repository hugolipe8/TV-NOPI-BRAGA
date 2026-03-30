"""
TV Nopi Braga — Servidor local de dados
Lê o Excel da Dropbox, parseia a folha RC e serve os dados em http://localhost:8765/dados
"""

import io
import json
import threading
import time
from datetime import datetime
from http.server import BaseHTTPRequestHandler, HTTPServer

import openpyxl
import requests

# ────────────────────────────────────────────────────────────────────────────────
# CONFIG
# ────────────────────────────────────────────────────────────────────────────────
EXCEL_URL = (
    "https://www.dropbox.com/scl/fi/y4i9m6v4q8snd2m3qljoh/Motherboard-2026.xlsx"
    "?rlkey=4px2hpxbg8p6fot2l65bkdamg&st=4h2vu72e&dl=1"
)
PORT             = 8765
REFRESH_INTERVAL = 5 * 60  # segundos

MONTH_OFFSETS = [0, 17, 30, 42, 54, 66, 78, 90, 102, 114, 126, 138]
MONTH_NAMES   = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]

# Nomes de linhas a ignorar (subtotais, cabeçalhos, cessados)
SKIP_NAMES = {"brg", "cg", "ag", "fp", "cm"}
# Sentinelas que marcam o fim dos dados activos
STOP_NAMES = {"total geral", "cessados"}

# ────────────────────────────────────────────────────────────────────────────────
# CACHE THREAD-SAFE
# ────────────────────────────────────────────────────────────────────────────────
_cache: dict | None = None
_cache_lock = threading.Lock()


def to_int(val) -> int:
    """Converte um valor de célula Excel para int, retorna 0 se inválido."""
    if val is None:
        return 0
    try:
        return int(float(str(val)))
    except (ValueError, TypeError):
        return 0


def fetch_and_parse() -> None:
    """Descarrega o Excel e actualiza o cache global."""
    global _cache
    try:
        print(f"[{_ts()}] A descarregar Excel...", flush=True)
        resp = requests.get(EXCEL_URL, timeout=60)
        resp.raise_for_status()

        wb = openpyxl.load_workbook(
            io.BytesIO(resp.content), read_only=True, data_only=True
        )

        if "RC" not in wb.sheetnames:
            raise ValueError('Folha "RC" não encontrada no ficheiro.')

        ws   = wb["RC"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        now = datetime.now()
        mi  = now.month - 1          # 0-based (Janeiro = 0)
        off = MONTH_OFFSETS[mi]

        # ── Totais BRG (linha índice 66) ──────────────────────────────────────
        brg_row   = rows[66] if len(rows) > 66 else ()
        total_ang = to_int(brg_row[off + 2]) if len(brg_row) > off + 2 else 0
        total_cont = to_int(brg_row[off + 4]) if len(brg_row) > off + 4 else 0

        # ── Consultores (a partir da linha índice 67) ─────────────────────────
        consultores = []
        for row in rows[67:]:
            if len(row) <= off:
                continue
            name = str(row[off] or "").strip()
            if not name:
                continue
            name_lower = name.lower()
            if name_lower in STOP_NAMES:
                break
            if name_lower in SKIP_NAMES:
                continue
            ang  = to_int(row[off + 2]) if len(row) > off + 2 else 0
            cont = to_int(row[off + 4]) if len(row) > off + 4 else 0
            if ang > 0 or cont > 0:
                consultores.append({"nome": name, "ang": ang, "cont": cont})

        data = {
            "mes":        MONTH_NAMES[mi],
            "ano":        now.year,
            "totalAng":   total_ang,
            "totalCont":  total_cont,
            "consultores": consultores,
            "atualizado": now.strftime("%H:%M"),
        }

        with _cache_lock:
            _cache = data

        print(
            f"[{_ts()}] OK — {data['mes']} {data['ano']} | "
            f"Ang {total_ang} | Cont {total_cont} | "
            f"{len(consultores)} consultores",
            flush=True,
        )

    except Exception as exc:
        print(f"[{_ts()}] ERRO ao carregar dados: {exc}", flush=True)


def _ts() -> str:
    return datetime.now().strftime("%H:%M:%S")


def _refresh_loop() -> None:
    while True:
        fetch_and_parse()
        time.sleep(REFRESH_INTERVAL)


# ────────────────────────────────────────────────────────────────────────────────
# HTTP HANDLER
# ────────────────────────────────────────────────────────────────────────────────
class Handler(BaseHTTPRequestHandler):
    def do_OPTIONS(self):
        self._cors(200)
        self.end_headers()

    def do_GET(self):
        if self.path not in ("/dados", "/dados/"):
            self._cors(404)
            self.end_headers()
            return

        with _cache_lock:
            data = _cache

        if data is None:
            body = json.dumps(
                {"erro": "Dados ainda não disponíveis. Aguarde alguns segundos."},
                ensure_ascii=False,
            ).encode("utf-8")
            self._cors(503)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self._cors(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-cache")
        self.end_headers()
        self.wfile.write(body)

    def _cors(self, code: int):
        self.send_response(code)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def log_message(self, fmt, *args):
        # Silencioso — não mostra pedidos HTTP no terminal
        pass


# ────────────────────────────────────────────────────────────────────────────────
# MAIN
# ────────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print(f"[{_ts()}] TV Nopi Braga — Servidor de dados", flush=True)
    print(f"[{_ts()}] A iniciar em http://localhost:{PORT}/dados", flush=True)

    # Primeira leitura em thread para não bloquear o arranque do HTTP
    threading.Thread(target=_refresh_loop, daemon=True).start()

    try:
        server = HTTPServer(("localhost", PORT), Handler)
        print(f"[{_ts()}] Servidor pronto. CTRL+C para parar.", flush=True)
        server.serve_forever()
    except KeyboardInterrupt:
        print(f"\n[{_ts()}] Servidor terminado.", flush=True)
